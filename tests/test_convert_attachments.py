"""Tests for convert.Converter: dispatch, sanitizing, images, nested messages,
caching, and the LibreOffice subprocess handling.

Everything that only needs pypdf, Pillow or reportlab runs for real on the
host. WeasyPrint-backed paths are exercised through a stand-in ``render_pdf``
(the real one is covered by test_render.py inside the container), and the
real LibreOffice test skips unless ``soffice`` is on PATH.
"""

from __future__ import annotations

import dis
import io
import os
import shutil
import stat
import subprocess
import sys
import threading
import time
import warnings
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from pst2pdf import convert, render
from pst2pdf.convert import Converter, detect, sanitize_pdf
from pst2pdf.model import Attachment, MessageModel, Options

from conftest import PNG_BYTES, make_ics, make_message, write_eml

# ---------------------------------------------------------------------------
# Fixture helpers
# ---------------------------------------------------------------------------


def _stub_pdf(path: Path, pages: int = 1, text: str = "stub") -> Path:
    c = canvas.Canvas(str(path), pagesize=letter)
    for i in range(pages):
        c.drawString(72, 700, f"{text} {i + 1}")
        c.showPage()
    c.save()
    return path


def _simple_pdf_bytes(pages: int = 2) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for i in range(pages):
        c.drawString(72, 700, f"page {i + 1}")
        c.showPage()
    c.save()
    return buf.getvalue()


def _nasty_pdf_bytes() -> bytes:
    """A PDF with a filled form field, a link, page JavaScript, document
    JavaScript, an OpenAction and an embedded file."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(72, 700, "Hello form")
    c.acroForm.textfield(name="fld", value="FILLED VALUE", x=72, y=600, width=200, height=20)
    c.linkURL("https://example.invalid/track", (72, 500, 200, 520))
    c.showPage()
    c.drawString(72, 700, "page two")
    c.showPage()
    c.save()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(buf.getvalue())))
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", DeprecationWarning)
        writer.add_js("app.alert('doc');")
    writer.add_attachment("secret.txt", b"embedded secret payload")
    js = DictionaryObject(
        {NameObject("/S"): NameObject("/JavaScript"), NameObject("/JS"): TextStringObject("app.alert('page')")}
    )
    writer.pages[0][NameObject("/AA")] = DictionaryObject({NameObject("/O"): js})
    writer.root_object[NameObject("/OpenAction")] = js
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


CHANNEL_MARKERS = (
    b"PAGE-XMP-SECRET",
    b"XOBJECT-XMP-SECRET",
    b"EMBEDDED-AF-SECRET",
    b"THUMB-SECRET",
)
#: Page keys the recursive ``excluded_keys`` removes (sanitizer and assembler).
EXCLUDED_CHANNEL_KEYS = ("/AF", "/Metadata", "/PieceInfo", "/Thumb", "/B", "/StructParents")
#: Plus the unknown keys only the sanitizer's page allowlist catches.
CHANNEL_KEYS = (*EXCLUDED_CHANNEL_KEYS, "/AAPL:PPK", "/Trans")


def channel_pdf_bytes() -> bytes:
    """A PDF whose page carries every non-content channel the sanitizer must
    close: an associated file with an embedded stream, XMP on the page and
    on its image XObject, PieceInfo, a thumbnail, article beads, a
    structure-tree link, and Apple's vendor key."""
    from pypdf.generic import ArrayObject, NumberObject, StreamObject
    from reportlab.lib.utils import ImageReader

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawImage(ImageReader(Image.new("RGB", (20, 20), (255, 0, 0))), 72, 600, 100, 100)
    c.drawString(72, 700, "visible text")
    c.showPage()
    c.save()
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(buf.getvalue())))
    page = writer.pages[0]

    def stream(data: bytes, **keys):
        obj = StreamObject()
        obj.set_data(data)
        for k, v in keys.items():
            obj[NameObject(k)] = v
        return writer._add_object(obj)

    page[NameObject("/Metadata")] = stream(
        b"<x:xmpmeta>PAGE-XMP-SECRET</x:xmpmeta>",
        **{"/Type": NameObject("/Metadata"), "/Subtype": NameObject("/XML")},
    )
    for ref in page["/Resources"]["/XObject"].values():
        ref.get_object()[NameObject("/Metadata")] = stream(b"<x:xmpmeta>XOBJECT-XMP-SECRET</x:xmpmeta>")
    embedded = stream(b"EMBEDDED-AF-SECRET", **{"/Type": NameObject("/EmbeddedFile")})
    filespec = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Filespec"),
                NameObject("/F"): TextStringObject("af.txt"),
                NameObject("/EF"): DictionaryObject({NameObject("/F"): embedded}),
            }
        )
    )
    page[NameObject("/AF")] = ArrayObject([filespec])
    page[NameObject("/PieceInfo")] = DictionaryObject(
        {NameObject("/Vendor"): DictionaryObject({NameObject("/Private"): TextStringObject("piece")})}
    )
    page[NameObject("/Thumb")] = stream(b"THUMB-SECRET", **{"/Width": NumberObject(1)})
    page[NameObject("/B")] = ArrayObject([])
    page[NameObject("/StructParents")] = NumberObject(0)
    page[NameObject("/AAPL:PPK")] = TextStringObject("apple private")
    page[NameObject("/Trans")] = DictionaryObject({NameObject("/S"): NameObject("/Wipe")})
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def assert_no_page_channels(path: Path, keys=CHANNEL_KEYS) -> None:
    """Every page and XObject in ``path`` is free of the channel keys and markers."""
    reader = PdfReader(str(path))
    for page in reader.pages:
        for key in keys:
            assert key not in page, key
        xobjects = page["/Resources"].get("/XObject") or {}
        for ref in xobjects.values():
            assert "/Metadata" not in ref.get_object()
    raw = path.read_bytes()
    for marker in CHANNEL_MARKERS:
        assert marker not in raw, marker
    assert b"/EmbeddedFile" not in raw


def _encrypted_pdf_bytes(user: str, owner: str = "owner") -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_simple_pdf_bytes(1))))
    writer.encrypt(user, owner)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def _png_bytes(size=(400, 150), mode="RGB", color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, "PNG")
    return buf.getvalue()


def _gif_bytes(frames: int) -> bytes:
    colors = [(255, 0, 0), (0, 255, 0), (0, 0, 255), (0, 0, 0)]
    images = [Image.new("RGB", (50, 80), colors[i % 4]) for i in range(frames)]
    buf = io.BytesIO()
    images[0].save(buf, "GIF", save_all=True, append_images=images[1:], duration=100)
    return buf.getvalue()


def _ooxml_zip(entries: dict[str, str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def minimal_docx(text: str = "Hello from a docx") -> bytes:
    return _ooxml_zip(
        {
            "[Content_Types].xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                "</Types>"
            ),
            "_rels/.rels": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                "</Relationships>"
            ),
            "word/document.xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"
            ),
        }
    )


def minimal_xlsx(cell: str = "Hello from a sheet") -> bytes:
    return _ooxml_zip(
        {
            "[Content_Types].xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                "</Types>"
            ),
            "_rels/.rels": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                "</Relationships>"
            ),
            "xl/workbook.xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
            ),
            "xl/_rels/workbook.xml.rels": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                "</Relationships>"
            ),
            "xl/worksheets/sheet1.xml": (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                f'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>{cell}</t></is></c></row></sheetData>'
                "</worksheet>"
            ),
        }
    )


def _plain_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("notes/a.txt", "hello " * 100)
        zf.writestr("b.doc", b"x" * 2048)
    return buf.getvalue()


@pytest.fixture
def options() -> Options:
    return Options(request_id="TEST", jobs=3, attachment_timeout=2)


@pytest.fixture
def stub_render_pdf(monkeypatch):
    """Replace WeasyPrint's render_pdf with a reportlab stand-in that records HTML."""
    calls: list[str] = []

    def fake(html: str, out_path: Path) -> render.Rendered:
        calls.append(html)
        _stub_pdf(Path(out_path))
        return render.Rendered()

    monkeypatch.setattr(render, "render_pdf", fake)
    monkeypatch.setattr(convert, "RENDER_IN_PROCESSES", False)  # a worker process would not see the stub
    return calls


@pytest.fixture
def converter(tmp_path, options):
    rendered: list[MessageModel] = []

    def render_message(model: MessageModel) -> Path:
        rendered.append(model)
        return _stub_pdf(tmp_path / f"msg-{len(rendered)}.pdf", pages=2)

    conv = Converter(tmp_path / "work", options, render_message)
    conv.rendered = rendered  # type: ignore[attr-defined]
    yield conv
    conv.close()


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "filename,ctype,expected",
    [
        ("Report.DOCX", "application/octet-stream", ("office", "docx")),
        ("a.xlsm", "", ("office", "xlsm")),
        ("scan.tif", "", ("image", "tif")),
        ("x.eml", "", ("eml", "eml")),
        ("noext", "application/pdf", ("pdf", "pdf")),
        ("noext", "text/html; charset=utf-8", ("html", "html")),
        ("noext", "application/vnd.ms-outlook", ("msg", "msg")),
        ("noext", "IMAGE/JPEG", ("image", "jpg")),
        ("weird.bin", "application/octet-stream", (None, "bin")),
    ],
)
def test_detect_by_extension_then_content_type(filename, ctype, expected):
    assert detect(filename, ctype, b"\x00" * 16) == expected


@pytest.mark.parametrize(
    "content,expected",
    [
        (b"%PDF-1.4 ...", ("pdf", "pdf")),
        (PNG_BYTES, ("image", "png")),
        (b"\xff\xd8\xff\xe0JFIF", ("image", "jpg")),
        (b"GIF89a....", ("image", "gif")),
        (b"RIFF\x00\x00\x00\x00WEBPVP8 ", ("image", "webp")),
        (b"{\\rtf1\\ansi hi}", ("office", "rtf")),
        (b"  BEGIN:VCALENDAR\r\nVERSION:2.0", ("ics", "ics")),
        (b"<!DOCTYPE html><html>", ("html", "html")),
        (b"\x00\x01\x02", (None, "")),
    ],
)
def test_detect_by_magic(content, expected):
    assert detect("blob", "application/octet-stream", content) == expected


def test_detect_zip_containers():
    assert detect("blob", "", minimal_docx()) == ("office", "docx")
    assert detect("blob", "", minimal_xlsx()) == ("office", "xlsx")
    odt = _ooxml_zip({"mimetype": "application/vnd.oasis.opendocument.text", "content.xml": "<x/>"})
    assert detect("blob", "", odt) == ("office", "odt")
    assert detect("blob", "", _plain_zip()) == ("zip", "zip")


# ---------------------------------------------------------------------------
# Placeholders
# ---------------------------------------------------------------------------


def test_empty_attachment_is_placeholder(converter):
    res = converter.convert(Attachment("empty.pdf", b"", "application/pdf"))
    assert (res.status, res.method, res.pdf_path, res.page_count) == ("placeholder", "none", None, 0)
    assert "empty" in res.reason


def test_oversize_attachment_is_placeholder(tmp_path):
    opts = Options(request_id="T", attachment_max_bytes=100, jobs=1)
    with Converter(tmp_path / "w", opts, lambda m: None) as conv:
        res = conv.convert(Attachment("big.png", _png_bytes(), "image/png"))
    assert res.status == "placeholder"
    assert "limit" in res.reason


def test_unknown_type_is_placeholder(converter):
    res = converter.convert(Attachment("thing.xyz", b"\x00\x01\x02\x03", "application/octet-stream"))
    assert res.status == "placeholder"
    assert ".xyz" in res.reason


def test_zip_is_listed_not_extracted(converter):
    res = converter.convert(Attachment("archive.zip", _plain_zip(), "application/zip"))
    assert res.status == "placeholder"
    lines = res.reason.splitlines()
    assert lines[0].startswith(("zip archive not expanded: 2 entries", "it is a zip archive of 2 entries"))
    assert "notes/a.txt (600 B)" in lines
    assert "b.doc (2.0 KB)" in lines
    # Nothing was extracted anywhere under the work dir.
    assert not list((converter.work_dir).rglob("a.txt"))


def test_bad_zip_is_placeholder(converter):
    res = converter.convert(Attachment("broken.zip", b"PK\x03\x04garbage", "application/zip"))
    assert res.status == "placeholder"
    assert "zip" in res.reason


def test_any_exception_becomes_placeholder(converter, monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("kaboom")

    monkeypatch.setattr(converter, "_via_image", boom)
    res = converter.convert(Attachment("x.png", _png_bytes(), "image/png"))
    assert res.status == "placeholder"
    assert "RuntimeError" in res.reason and "kaboom" in res.reason


# ---------------------------------------------------------------------------
# PDF pass-through
# ---------------------------------------------------------------------------


def test_pdf_passthrough_strips_active_content(converter):
    nasty = _nasty_pdf_bytes()
    before = PdfReader(io.BytesIO(nasty))
    assert list(before.attachments) == ["secret.txt"]
    assert "/AcroForm" in before.root_object and "/Names" in before.root_object

    res = converter.convert(Attachment("form.pdf", nasty, "application/pdf"))
    assert (res.status, res.method, res.page_count) == ("passthrough", "pdf", 2)
    assert res.pdf_path.parent == converter.out_dir

    after = PdfReader(str(res.pdf_path))
    root = after.root_object
    for key in ("/AcroForm", "/Names", "/OpenAction", "/AA"):
        assert key not in root
    assert list(after.attachments) == []
    for page in after.pages:
        assert "/Annots" not in page and "/AA" not in page
    raw = res.pdf_path.read_bytes()
    assert b"JavaScript" not in raw
    assert b"embedded secret payload" not in raw
    assert b"example.invalid" not in raw
    # The form field's value survived as page content (flattened).
    assert "FILLED VALUE" in after.pages[0].extract_text()
    assert "page two" in after.pages[1].extract_text()


def test_pdf_password_protected_is_placeholder(converter):
    res = converter.convert(Attachment("locked.pdf", _encrypted_pdf_bytes("user"), "application/pdf"))
    assert res.status == "placeholder"
    assert res.reason == "password protected"


def test_pdf_owner_password_only_passes_through(converter):
    res = converter.convert(Attachment("restricted.pdf", _encrypted_pdf_bytes(""), "application/pdf"))
    assert res.status == "passthrough"
    assert res.page_count == 1
    assert not PdfReader(str(res.pdf_path)).is_encrypted


def test_pdf_garbage_is_placeholder(converter):
    res = converter.convert(Attachment("junk.pdf", b"%PDF-1.7 this is not really a pdf", "application/pdf"))
    assert res.status == "placeholder"
    assert res.reason


def test_sanitize_pdf_reports_counts(tmp_path):
    stats = sanitize_pdf(_nasty_pdf_bytes(), tmp_path / "clean.pdf")
    assert stats == {"pages": 2, "flattened": 1, "annotations": 2}


def test_pdf_passthrough_strips_page_level_channels(converter):
    src = channel_pdf_bytes()
    before = PdfReader(io.BytesIO(src)).pages[0]
    for key in CHANNEL_KEYS:
        assert key in before, key
    for marker in CHANNEL_MARKERS:
        assert marker in src, marker

    res = converter.convert(Attachment("apple.pdf", src, "application/pdf"))
    assert (res.status, res.page_count) == ("passthrough", 1)
    assert_no_page_channels(res.pdf_path)
    after = PdfReader(str(res.pdf_path)).pages[0]
    assert set(after.keys()) <= convert.PAGE_ALLOWED_KEYS
    assert "visible text" in after.extract_text()
    assert len(after["/Resources"]["/XObject"]) == 1  # the image itself survives


def test_converted_pdf_over_page_limit_is_placeholder(tmp_path):
    opts = Options(request_id="T", jobs=1, attachment_max_pages=2)
    with Converter(tmp_path / "w", opts, lambda m: None) as conv:
        ok = conv.convert(Attachment("two.pdf", _simple_pdf_bytes(2), "application/pdf"))
        over = conv.convert(Attachment("three.pdf", _simple_pdf_bytes(3), "application/pdf"))
    assert ok.status == "passthrough"
    assert (over.status, over.pdf_path, over.page_count) == ("placeholder", None, 0)
    assert over.reason == "3 pages, more than the 2-page limit"
    assert list((tmp_path / "w" / "converted").iterdir()) == [ok.pdf_path]


def test_page_limit_default_is_two_thousand():
    assert Options(request_id="T").attachment_max_pages == 2000


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------


def test_wide_image_lands_on_landscape_letter(converter):
    res = converter.convert(Attachment("wide.png", _png_bytes((400, 150)), "image/png"))
    assert (res.status, res.method, res.page_count) == ("converted", "image", 1)
    box = PdfReader(str(res.pdf_path)).pages[0].mediabox
    assert (float(box.width), float(box.height)) == (792.0, 612.0)


def test_tall_image_lands_on_portrait_letter(converter):
    res = converter.convert(Attachment("tall.jpg", _jpeg_bytes((300, 900)), "image/jpeg"))
    assert res.page_count == 1
    box = PdfReader(str(res.pdf_path)).pages[0].mediabox
    assert (float(box.width), float(box.height)) == (612.0, 792.0)


def _jpeg_bytes(size):
    buf = io.BytesIO()
    Image.new("RGB", size, (10, 120, 200)).save(buf, "JPEG")
    return buf.getvalue()


def test_animated_gif_gets_one_page_per_frame(converter):
    res = converter.convert(Attachment("anim.gif", _gif_bytes(3), "image/gif"))
    assert (res.status, res.page_count) == ("converted", 3)


def test_transparent_and_palette_images_convert(converter):
    rgba = _png_bytes((60, 40), "RGBA", (255, 0, 0, 128))
    assert converter.convert(Attachment("t.png", rgba, "image/png")).status == "converted"
    assert converter.convert(Attachment("one.png", PNG_BYTES, "image/png")).status == "converted"


def test_corrupt_image_is_placeholder(converter):
    res = converter.convert(Attachment("bad.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 40, "image/png"))
    assert res.status == "placeholder"


def test_images_module_caps_pillow_pixels():
    from pst2pdf import images

    assert images.MAX_IMAGE_PIXELS == 50_000_000
    assert Image.MAX_IMAGE_PIXELS == images.MAX_IMAGE_PIXELS


@pytest.mark.filterwarnings("ignore::PIL.Image.DecompressionBombWarning")
def test_image_over_pixel_limit_is_placeholder(converter, monkeypatch):
    # 100x100 = 10,000 pixels: over the (lowered) limit but under Pillow's own
    # bomb threshold of twice the limit, so the converter's check is what fires.
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 8_000)
    res = converter.convert(Attachment("big.png", _png_bytes((100, 100)), "image/png"))
    assert (res.status, res.method, res.pdf_path) == ("placeholder", "none", None)
    assert res.reason == "image is 100x100 px (10,000 pixels), over the 8,000-pixel limit"


def test_decompression_bomb_is_placeholder(converter, monkeypatch):
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 10)  # 100x100 is > 2x: Pillow refuses to open
    res = converter.convert(Attachment("bomb.png", _png_bytes((100, 100)), "image/png"))
    assert res.status == "placeholder"
    assert res.reason.startswith("image is too large to open safely")
    assert "decompression bomb" in res.reason


# ---------------------------------------------------------------------------
# Text (reportlab) and HTML (WeasyPrint stand-in)
# ---------------------------------------------------------------------------


def _pdf_text(res) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(str(res.pdf_path)).pages)


def test_text_is_written_directly_as_real_text(converter, stub_render_pdf):
    res = converter.convert(Attachment("notes.txt", b"\xef\xbb\xbfa < b & c\nline2", "text/plain"))
    assert (res.status, res.method, res.page_count) == ("converted", "text", 1)
    assert stub_render_pdf == []  # no WeasyPrint
    text = _pdf_text(res)
    assert "a < b & c" in text and "line2" in text
    assert "\ufeff" not in text  # BOM stripped
    page = PdfReader(str(res.pdf_path)).pages[0]
    assert [float(v) for v in page.mediabox] == [0, 0, 612, 792]


def test_text_decodes_cp1252_fallback(converter, stub_render_pdf):
    res = converter.convert(Attachment("latin.csv", b"caf\xe9,\x93quoted\x94", "text/csv"))
    assert res.status == "converted"
    text = _pdf_text(res)
    assert "café" in text
    assert "“quoted”" in text


def test_text_the_font_cannot_draw_goes_to_weasyprint(converter, stub_render_pdf, monkeypatch):
    monkeypatch.setattr(convert.textpdf, "covers", lambda text: False)
    res = converter.convert(Attachment("notes.txt", b"a < b", "text/plain"))
    assert (res.status, res.method) == ("converted", "text")
    assert "<pre>a &lt; b</pre>" in stub_render_pdf[-1]
    assert "size: Letter" in stub_render_pdf[-1]


def test_long_text_paginates_at_59_lines(converter, stub_render_pdf):
    body = "".join(f"2026-09-01 INFO worker processed record {i}\n" for i in range(59 * 3 + 1)).encode()
    res = converter.convert(Attachment("sync.log", body, "text/plain"))
    assert (res.status, res.method, res.page_count) == ("converted", "text", 4)
    pages = PdfReader(str(res.pdf_path)).pages
    assert "record 0" in pages[0].extract_text() and "record 58" in pages[0].extract_text()
    assert "record 59" in pages[1].extract_text()
    assert "record 177" in pages[3].extract_text()


def test_oversize_text_is_placeholder(converter, stub_render_pdf, monkeypatch):
    monkeypatch.setattr(convert, "TEXT_MAX_BYTES", 10)
    res = converter.convert(Attachment("huge.log", b"x" * 11, "text/plain"))
    assert res.status == "placeholder"
    assert stub_render_pdf == []


def test_html_goes_through_hygiene(converter, stub_render_pdf):
    src = (
        b"<html><head><title>t</title><script>evil()</script></head>"
        b'<body><p>Hi <a href="https://example.invalid/x">link</a></p>'
        b'<img src="https://example.invalid/pixel.gif" alt="tracker"></body></html>'
    )
    res = converter.convert(Attachment("page.html", src, "text/html"))
    assert (res.status, res.method) == ("converted", "html")
    html = stub_render_pdf[-1]
    assert "<script>" not in html and "<title>" not in html
    assert "href=" not in html
    assert "Hi" in html


def test_blocked_remote_resources_in_html_attachments_are_counted(converter, monkeypatch):
    """The legacy layout read a module global for this; the volumes read the
    converter's own count, which the run adds to the message pages'."""
    monkeypatch.setattr(convert, "RENDER_IN_PROCESSES", False)

    def fake(html: str, out_path: Path) -> render.Rendered:
        _stub_pdf(Path(out_path))
        return render.Rendered(blocked=2)

    monkeypatch.setattr(render, "render_pdf", fake)
    converter.convert(Attachment("a.html", b"<p>one</p>", "text/html"))
    converter.convert(Attachment("b.html", b"<p>two</p>", "text/html"))
    assert converter.blocked_urls == 4


def test_html_falls_back_to_source_when_cleaner_missing(converter, stub_render_pdf, monkeypatch):
    monkeypatch.setattr(convert, "_clean_html", lambda fragment: None)
    res = converter.convert(Attachment("page.htm", b"<p>x</p>", "text/html"))
    assert res.status == "converted"
    assert "<pre>&lt;p&gt;x&lt;/p&gt;</pre>" in stub_render_pdf[-1]


# ---------------------------------------------------------------------------
# Nested messages
# ---------------------------------------------------------------------------


def test_eml_attachment_rendered_through_message_template(converter, tmp_path):
    nested = make_message(subject="Nested subject", attachments=[("inner.docx", b"PK\x03\x04")])
    eml_path = tmp_path / "nested.eml"
    write_eml(nested, eml_path)

    res = converter.convert(Attachment("fwd.eml", eml_path.read_bytes(), "message/rfc822"))
    assert (res.status, res.method, res.page_count) == ("converted", "message", 2)
    assert res.pdf_path.parent == converter.out_dir
    model = converter.rendered[-1]
    assert model.subject == "Nested subject"
    assert [a.filename for a in model.attachments] == ["inner.docx"]
    # The nested message's own attachment is listed but not converted.
    assert len(list(converter.out_dir.iterdir())) == 1


class _FakeAttachment:
    def __init__(self, name, data, mimetype="application/octet-stream", cid=None):
        self.longFilename = name
        self.name = name
        self.shortFilename = name
        self.data = data
        self.mimetype = mimetype
        self.cid = cid
        self.contentId = cid


class _FakeEmbedded:
    def exportBytes(self):
        return b"\xd0\xcf\x11\xe0inner"


class _FakeMsg:
    subject = "From an Outlook item"
    sender = "Carol <carol@example.com>"
    to = "dave@example.com"
    cc = None
    date = None
    messageId = "<abc@example.com>"
    body = "plain body"
    htmlBody = (
        b'<html><body><p>See <img src="cid:img1"> and <a href="https://x.invalid">this</a></p></body></html>'
    )
    closed = False

    def __init__(self, path):
        from datetime import datetime, timezone

        self.date = datetime(2024, 3, 15, 9, 0, tzinfo=timezone.utc)
        self.attachments = [
            _FakeAttachment("photo.png", PNG_BYTES, "image/png", cid="img1"),
            _FakeAttachment("spreadsheet.xlsx", b"PK\x03\x04data"),
            _FakeAttachment("Embedded message", _FakeEmbedded()),
        ]

    def close(self):
        _FakeMsg.closed = True


def test_msg_attachment_mapped_to_model(converter, monkeypatch):
    import extract_msg

    monkeypatch.setattr(extract_msg, "openMsg", lambda path, **kw: _FakeMsg(path))
    res = converter.convert(Attachment("item.msg", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "application/vnd.ms-outlook"))
    assert (res.status, res.method, res.page_count) == ("converted", "message", 2)
    model = converter.rendered[-1]
    assert model.subject == "From an Outlook item"
    assert model.from_addr == "Carol <carol@example.com>"
    assert model.to_addrs == "dave@example.com" and model.cc_addrs == ""
    assert model.date is not None and model.date.year == 2024
    assert model.message_id == "abc@example.com"
    assert "data:image/png;base64," in model.body_html and "cid:" not in model.body_html
    assert model.has_images
    assert [a.filename for a in model.attachments] == ["spreadsheet.xlsx", "Embedded message.msg"]
    assert model.attachments[1].content.startswith(b"\xd0\xcf\x11\xe0")
    assert _FakeMsg.closed


def test_ics_attachment_uses_calendar_parser(converter):
    pytest.importorskip("pst2pdf.calendar")
    res = converter.convert(Attachment("invite.ics", make_ics(summary="Budget meeting"), "text/calendar"))
    assert (res.status, res.method, res.page_count) == ("converted", "calendar", 2)
    model = converter.rendered[-1]
    assert model.kind == "calendar"
    assert model.subject == "Budget meeting"


def test_ics_falls_back_to_text_without_calendar_module(converter, stub_render_pdf, monkeypatch):
    monkeypatch.setitem(sys.modules, "pst2pdf.calendar", None)
    res = converter.convert(Attachment("invite.ics", make_ics(), "text/calendar"))
    assert (res.status, res.method) == ("converted", "text")
    assert "BEGIN:VCALENDAR" in _pdf_text(res)
    assert converter.rendered == []


# ---------------------------------------------------------------------------
# Cache and concurrency
# ---------------------------------------------------------------------------


def test_identical_bytes_converted_once_across_threads(tmp_path, options, monkeypatch):
    calls = []
    lock = threading.Lock()
    with Converter(tmp_path / "w", options, lambda m: None) as conv:
        original = conv._convert_uncached

        def counting(att, digest):
            with lock:
                calls.append(digest)
            return original(att, digest)

        monkeypatch.setattr(conv, "_convert_uncached", counting)
        png = _png_bytes()
        atts = [Attachment(f"copy{i}.png", png, "image/png") for i in range(40)]
        atts.append(Attachment("other.png", _png_bytes((10, 10)), "image/png"))

        results: list[list] = [[] for _ in range(4)]

        def worker(slot):
            results[slot] = conv.convert_all(atts)

        threads = [threading.Thread(target=worker, args=(i,)) for i in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

    assert len(calls) == 2
    for batch in results:
        assert len(batch) == 41
        assert all(r.status == "converted" for r in batch)
        assert len({r.pdf_path for r in batch[:40]}) == 1
        assert batch[40].pdf_path != batch[0].pdf_path
    assert results[0][0] is results[3][7]


def test_convert_all_preserves_order(converter):
    atts = [
        Attachment("a.png", _png_bytes((20, 20)), "image/png"),
        Attachment("b.zip", _plain_zip(), "application/zip"),
        Attachment("c.pdf", _simple_pdf_bytes(3), "application/pdf"),
    ]
    out = converter.convert_all(atts)
    assert [r.attachment.filename for r in out] == ["a.png", "b.zip", "c.pdf"]
    assert [r.method for r in out] == ["image", "none", "pdf"]
    assert [r.page_count for r in out] == [1, 0, 3]


def test_never_writes_outside_work_dir(tmp_path, options):
    work = tmp_path / "work"
    with Converter(work, options, lambda m: None) as conv:
        conv.convert(Attachment("a.png", _png_bytes(), "image/png"))
        conv.convert(Attachment("c.pdf", _simple_pdf_bytes(), "application/pdf"))
        conv.convert(Attachment("z.zip", _plain_zip(), "application/zip"))
        assert (work / "convert-tmp").is_dir()
        assert list((work / "convert-tmp").iterdir()) == []  # per-call dirs cleaned
    assert sorted(p.name for p in tmp_path.iterdir()) == ["work"]
    assert sorted(p.name for p in work.iterdir()) == ["converted"]
    names = [p.name for p in (work / "converted").iterdir()]
    assert len(names) == 2 and all(len(n) == 68 and n.endswith(".pdf") for n in names)


# ---------------------------------------------------------------------------
# LibreOffice subprocess handling
# ---------------------------------------------------------------------------


def test_missing_soffice_is_placeholder(converter):
    converter.soffice = None
    res = converter.convert(Attachment("doc.docx", minimal_docx(), ""))
    assert res.status == "placeholder"
    assert "LibreOffice not available" == res.reason


def _fake_soffice(tmp_path: Path, body: str) -> Path:
    script = tmp_path / "fake-soffice.sh"
    script.write_text("#!/bin/sh\n" + body)
    script.chmod(script.stat().st_mode | stat.S_IXUSR)
    return script


@pytest.mark.skipif(os.name != "posix", reason="process groups")
def test_soffice_timeout_kills_process_group(converter, tmp_path):
    pid_file = tmp_path / "grandchild.pid"
    # Like oosplash forking soffice.bin: the grandchild outlives the child.
    converter.soffice = str(
        _fake_soffice(tmp_path, f"sleep 300 &\necho $! > '{pid_file}'\nwait\n")
    )
    res = converter.convert(Attachment("slow.docx", minimal_docx(), ""))
    assert res.status == "placeholder"
    assert res.reason == "conversion timed out after 2 s"
    grandchild = int(pid_file.read_text())
    assert _dead_within(grandchild, 5), f"grandchild {grandchild} survived the kill"


def _dead_within(pid: int, seconds: float) -> bool:
    """True once pid is gone or a zombie. The killed grandchild is an orphan,
    and until its new parent reaps it, os.kill(pid, 0) still succeeds. In the
    container that parent is tini, which under full-suite load may take a
    moment; on a macOS host launchd reaps it at once."""
    deadline = time.monotonic() + seconds
    while True:
        try:
            os.kill(pid, 0)
        except ProcessLookupError:
            return True
        try:
            stat_line = Path(f"/proc/{pid}/stat").read_text()
        except OSError:
            pass  # no /proc (macOS), or reaped since the kill() above
        else:
            # The state follows the parenthesised command name, which may
            # itself contain ") ".
            if stat_line.rpartition(")")[2].split()[0] == "Z":
                return True
        if time.monotonic() >= deadline:
            return False
        time.sleep(0.05)


@pytest.mark.skipif(os.name != "posix", reason="POSIX signals")
def test_dead_within_rejects_a_live_process():
    live = subprocess.Popen(["sleep", "30"])
    try:
        assert not _dead_within(live.pid, 0.2)
    finally:
        live.kill()
        live.wait()
    assert _dead_within(live.pid, 0.2)


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="/proc state")
def test_dead_within_accepts_a_zombie():
    child = subprocess.Popen(["sleep", "30"])
    child.kill()
    try:
        assert _dead_within(child.pid, 5)
    finally:
        child.wait()


@pytest.mark.skipif(os.name != "posix", reason="POSIX fork semantics")
def test_preexec_hook_does_nothing_that_could_deadlock_the_child():
    """The RLIMIT_AS hook runs between fork and exec, in a child forked from a
    multi-threaded parent: a lock another thread held at fork time is held
    forever there. Anything but bare syscalls -- an import above all -- risks
    hanging the child until the conversion timeout kills it."""
    assert convert.resource is not None
    opnames = {i.opname for i in dis.get_instructions(convert._limit_address_space)}
    assert "IMPORT_NAME" not in opnames and "IMPORT_FROM" not in opnames, opnames


@pytest.mark.skipif(os.name != "posix", reason="shell script stand-in")
def test_soffice_invocation_is_isolated(converter, tmp_path):
    """The fake soffice records its environment and writes out/input.pdf."""
    stub = _stub_pdf(tmp_path / "stub.pdf")
    record = tmp_path / "record.txt"
    converter.soffice = str(
        _fake_soffice(
            tmp_path,
            f"echo \"$@\" > '{record}'\n"
            f"echo \"cwd=$(pwd)\" >> '{record}'\n"
            f"echo \"home=$HOME\" >> '{record}'\n"
            f"echo \"tmpdir=$TMPDIR\" >> '{record}'\n"
            # copy the profile registry so we can inspect what soffice would read
            "for a in \"$@\"; do case \"$a\" in -env:UserInstallation=file://*) "
            f"cp \"${{a#-env:UserInstallation=file://}}/user/registrymodifications.xcu\" '{tmp_path}/xcu.xml';; esac; done\n"
            # find --outdir and produce the expected output
            "while [ $# -gt 0 ]; do if [ \"$1\" = --outdir ]; then out=\"$2\"; fi; shift; done\n"
            f"cp '{stub}' \"$out/input.pdf\"\n",
        )
    )
    res = converter.convert(Attachment("memo.doc", b"\xd0\xcf\x11\xe0" + b"\x00" * 32, "application/msword"))
    assert (res.status, res.method, res.page_count) == ("converted", "libreoffice", 1)

    lines = record.read_text().splitlines()
    argv = lines[0]
    for flag in ("--headless", "--norestore", "--nolockcheck", "--nologo", "--convert-to pdf"):
        assert flag in argv
    assert argv.endswith("/input.doc")
    work = str(converter.work_dir.resolve())
    assert f"-env:UserInstallation=file://{work}/convert-tmp/" in argv
    for line in lines[1:]:
        key, _, value = line.partition("=")
        assert Path(value).resolve().is_relative_to(converter.work_dir.resolve()), line

    xcu = (tmp_path / "xcu.xml").read_text()
    assert _xcu_values(xcu) == PROFILE_EXPECTED
    # per-call temp dir is gone afterwards
    assert list((converter.work_dir / "convert-tmp").iterdir()) == []


#: (registry path, property) -> value the generated profile must carry.
PROFILE_EXPECTED = {
    ("/org.openoffice.Office.Common/Security/Scripting", "MacroSecurityLevel"): "3",
    ("/org.openoffice.Office.Common/Security/Scripting", "DisableMacrosExecution"): "true",
    ("/org.openoffice.Office.Common/Security/Scripting", "HyperlinksWithCtrlClick"): "true",
    ("/org.openoffice.Office.Common/Security/Scripting", "BlockUntrustedRefererLinks"): "true",
    ("/org.openoffice.Office.Common/Misc", "FirstRun"): "false",
    ("/org.openoffice.Office.Common/Misc", "ShowTipOfTheDay"): "false",
    ("/org.openoffice.Office.Common/Save/Document", "AutoSave"): "false",
    ("/org.openoffice.Office.Writer/Content/Update", "Link"): "0",  # Writer: 0 = never
    ("/org.openoffice.Office.Calc/Content/Update", "Link"): "1",  # Calc: 1 = never (0 is always)
}


def _xcu_values(xcu: str) -> dict[tuple[str, str], str]:
    """(item path, prop name) -> value for every prop in a registrymodifications.xcu."""
    import re

    found = {}
    for path, body in re.findall(r'<item oor:path="([^"]+)">(.*?)</item>', xcu, re.DOTALL):
        for name, value in re.findall(r'<prop oor:name="([^"]+)"[^>]*>\s*<value>([^<]*)</value>', body):
            found[(path, name)] = value
    return found


def test_generated_profile_carries_every_security_setting(tmp_path):
    convert._write_profile(tmp_path / "profile")
    xcu = (tmp_path / "profile" / "user" / "registrymodifications.xcu").read_text()
    assert _xcu_values(xcu) == PROFILE_EXPECTED


# ---------------------------------------------------------------------------
# Spreadsheet pre-processing
# ---------------------------------------------------------------------------


def _wide_workbook(columns: int = 40, rows: int = 30, tall_rows=(2, 5, 9)) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Wide"
    for r in range(1, rows + 1):
        for c in range(1, columns + 1):
            ws.cell(row=r, column=c, value=f"r{r}c{c}")
    for r in tall_rows:
        ws.row_dimensions[r].height = 90
    narrow = wb.create_sheet("Narrow")
    for r in range(1, 5):
        for c in range(1, 4):
            narrow.cell(row=r, column=c, value=r * c)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_prepare_spreadsheet_fits_width_and_clears_row_heights():
    from openpyxl import load_workbook

    src = _wide_workbook()
    before = load_workbook(io.BytesIO(src))["Wide"]
    assert before.row_dimensions[2].customHeight and before.row_dimensions[2].height == 90
    assert before.page_setup.fitToWidth is None or before.page_setup.fitToWidth != 1

    fitted = convert.prepare_spreadsheet(src)
    assert fitted is not None and fitted != src
    wb = load_workbook(io.BytesIO(fitted))
    assert wb.sheetnames == ["Wide", "Narrow"]
    for ws in wb.worksheets:
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 0
        assert ws.page_setup.paperSize == 1  # Letter
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert not any(d.customHeight for d in ws.row_dimensions.values())
    assert wb["Wide"].page_setup.orientation == "landscape"
    assert wb["Narrow"].page_setup.orientation != "landscape"
    # Cell content is untouched.
    assert wb["Wide"]["AN30"].value == "r30c40"
    assert wb["Narrow"]["C4"].value == 12


def test_prepare_spreadsheet_falls_back_on_unreadable_input():
    assert convert.prepare_spreadsheet(b"not a workbook") is None
    assert convert.prepare_spreadsheet(minimal_docx()) is None


def _workbook_with_unparsable_header() -> bytes:
    """A workbook whose printed header openpyxl cannot parse.

    Real archives carry these: openpyxl warns, drops the header, and reads
    the rest of the file fine.
    """
    src = zipfile.ZipFile(io.BytesIO(_wide_workbook(columns=3, rows=3, tall_rows=())))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.decode().replace(
                    "</worksheet>",
                    "<headerFooter><oddHeader>&amp;Zbogus&amp;</oddHeader>"
                    "</headerFooter></worksheet>",
                ).encode()
            zf.writestr(info, data)
    return out.getvalue()


def test_prepare_spreadsheet_logs_openpyxl_warnings_instead_of_printing_them(caplog):
    """openpyxl's warnings must not reach stderr: the wrapper prints anything
    that is not an INFO/WARNING log line, so a stray UserWarning lands in the
    operator's progress output looking like a crash."""
    payload = _workbook_with_unparsable_header()

    with caplog.at_level("DEBUG", logger="pst2pdf.convert"):
        with warnings.catch_warnings(record=True) as leaked:
            warnings.simplefilter("always")
            fitted = convert.prepare_spreadsheet(payload, "Q3 budget.xlsx")
            # A second workbook with the same defect is reported again, rather
            # than being swallowed by the once-per-site default filter.
            convert.prepare_spreadsheet(payload, "Headcount.xlsx")

    assert fitted is not None  # the rest of the workbook still converts
    assert [str(w.message) for w in leaked] == []
    logged = [r.getMessage() for r in caplog.records if "openpyxl dropped" in r.getMessage()]
    assert len(logged) == 2
    assert logged[0].startswith("Q3 budget.xlsx: ")
    assert logged[1].startswith("Headcount.xlsx: ")
    assert "header or footer" in logged[0]


@pytest.mark.skipif(os.name != "posix", reason="shell script stand-in")
def test_xlsx_is_preprocessed_before_soffice(converter, tmp_path):
    """The fake soffice keeps a copy of whatever it was asked to convert."""
    from openpyxl import load_workbook

    stub = _stub_pdf(tmp_path / "stub.pdf")
    keep = tmp_path / "staged"
    keep.mkdir()
    converter.soffice = str(
        _fake_soffice(
            tmp_path,
            "while [ $# -gt 1 ]; do if [ \"$1\" = --outdir ]; then out=\"$2\"; fi; shift; done\n"
            f"cp \"$1\" '{keep}/'\n"
            f"cp '{stub}' \"$out/input.pdf\"\n",
        )
    )
    res = converter.convert(Attachment("Budget.XLSM", _wide_workbook(), ""))
    assert (res.status, res.method) == ("converted", "libreoffice")
    (staged,) = keep.iterdir()
    assert staged.name == "input.xlsx"  # rewritten as a plain xlsx
    ws = load_workbook(str(staged))["Wide"]
    assert (ws.page_setup.fitToWidth, ws.page_setup.fitToHeight, ws.page_setup.orientation) == (1, 0, "landscape")

    # Legacy .xls and non-spreadsheets are staged as they came.
    for p in keep.iterdir():
        p.unlink()
    xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    res = converter.convert(Attachment("old.xls", xls, ""))
    assert res.status == "converted"
    (staged,) = keep.iterdir()
    assert staged.name == "input.xls" and staged.read_bytes() == xls


@pytest.mark.skipif(os.name != "posix", reason="shell script stand-in")
def test_broken_xlsx_is_handed_over_unchanged(converter, tmp_path):
    keep = tmp_path / "staged"
    keep.mkdir()
    stub = _stub_pdf(tmp_path / "stub.pdf")
    converter.soffice = str(
        _fake_soffice(
            tmp_path,
            "while [ $# -gt 1 ]; do if [ \"$1\" = --outdir ]; then out=\"$2\"; fi; shift; done\n"
            f"cp \"$1\" '{keep}/'\n"
            f"cp '{stub}' \"$out/input.pdf\"\n",
        )
    )
    payload = minimal_xlsx()  # valid to LibreOffice, but too bare for openpyxl
    if convert.prepare_spreadsheet(payload) is not None:
        payload = b"PK\x03\x04 truncated"
        assert convert.prepare_spreadsheet(payload) is None
    res = converter.convert(Attachment("sheet.xlsx", payload, ""))
    assert res.status == "converted"
    (staged,) = keep.iterdir()
    assert staged.name == "input.xlsx" and staged.read_bytes() == payload


@pytest.mark.skipif(os.name != "posix", reason="shell script stand-in")
def test_soffice_failure_is_placeholder(converter, tmp_path):
    converter.soffice = str(_fake_soffice(tmp_path, "echo 'Error: source file could not be loaded' >&2\nexit 1\n"))
    res = converter.convert(Attachment("memo.docx", minimal_docx(), ""))
    assert res.status == "placeholder"
    assert "could not be loaded" in res.reason


needs_soffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


@needs_soffice
@pytest.mark.parametrize("name,payload", [("memo.docx", minimal_docx()), ("sheet.xlsx", minimal_xlsx())])
def test_real_libreoffice_conversion(tmp_path, name, payload):
    opts = Options(request_id="T", jobs=2, attachment_timeout=120)
    with Converter(tmp_path / "w", opts, lambda m: None) as conv:
        res = conv.convert(Attachment(name, payload, ""))
    assert (res.status, res.method) == ("converted", "libreoffice"), res.reason
    assert res.page_count >= 1
    text = PdfReader(str(res.pdf_path)).pages[0].extract_text()
    assert "Hello from a" in text


@needs_soffice
def test_real_libreoffice_prints_wide_sheet_fitted_to_width(tmp_path):
    """40 columns of text fit one page across (landscape Letter) instead of
    being sliced into column strips, and the tall rows do not pad the sheet
    into extra pages."""
    opts = Options(request_id="T", jobs=1, attachment_timeout=180)
    with Converter(tmp_path / "w", opts, lambda m: None) as conv:
        res = conv.convert(Attachment("wide.xlsx", _wide_workbook(), ""))
    assert (res.status, res.method) == ("converted", "libreoffice"), res.reason
    reader = PdfReader(str(res.pdf_path))
    first = reader.pages[0]
    assert float(first.mediabox.width) > float(first.mediabox.height)  # landscape
    text = first.extract_text()
    assert "r1c1" in text and "r1c40" in text  # first and last column on one page
    assert len(reader.pages) <= 3  # the whole 30x40 sheet plus the Narrow sheet


@needs_soffice
def test_real_libreoffice_honours_generated_profile(tmp_path):
    """LibreOffice rewrites registrymodifications.xcu on exit with the items
    it parsed into its configuration layer; ours must survive that rewrite.

    (A behavioural check with a macro-bearing document is not possible:
    under --convert-to no document load event fires at any security level.)
    """
    import re

    profile = tmp_path / "profile"
    convert._write_profile(profile)
    xcu = profile / "user" / "registrymodifications.xcu"
    before = xcu.read_text()
    src = tmp_path / "input.docx"
    src.write_bytes(minimal_docx())
    (tmp_path / "out").mkdir()
    cmd = [
        shutil.which("soffice"), "--headless", "--norestore", "--nolockcheck", "--nologo",
        f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "pdf",
        "--outdir", str(tmp_path / "out"), str(src),
    ]
    env = {"HOME": str(tmp_path), "PATH": os.environ["PATH"]}
    result = convert._run_killable(cmd, cwd=tmp_path, env=env, timeout=120)
    assert result is not None and result[0] == 0, result
    after = xcu.read_text()
    assert after != before and "/org.openoffice.Setup" in after  # LibreOffice really rewrote it
    kept = _xcu_values(after)
    for key, value in PROFILE_EXPECTED.items():
        assert kept.get(key) == value, key


@needs_soffice
@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="/proc scan")
def test_real_libreoffice_timeout_leaves_no_processes(tmp_path):
    big = minimal_docx("lorem ipsum " * 40 + "</w:t></w:r></w:p><w:p><w:r><w:t>" * 4000)
    opts = Options(request_id="T", jobs=1, attachment_timeout=0.25)  # type: ignore[arg-type]
    with Converter(tmp_path / "w", opts, lambda m: None) as conv:
        res = conv.convert(Attachment("slow.docx", big, ""))
    assert res.status == "placeholder"
    assert res.reason == "conversion timed out after 0.25 s"
    stray = []
    for pid in os.listdir("/proc"):
        if pid.isdigit():
            try:
                cmd = Path(f"/proc/{pid}/cmdline").read_bytes()
            except OSError:
                continue
            if b"soffice" in cmd or b"oosplash" in cmd:
                stray.append(int(pid))
    # A SIGKILLed process keeps its cmdline until it has finished exiting,
    # which under full-suite load can lag the kill; give it the same grace
    # as test_soffice_timeout_kills_process_group.
    assert [pid for pid in stray if not _dead_within(pid, 5)] == []


# ---------------------------------------------------------------------------
# Layout in worker processes
# ---------------------------------------------------------------------------


def test_text_lays_out_in_worker_processes(tmp_path):
    """With jobs > 1, text is laid out by the worker pool (real processes,
    no stubs), several files at once, each with its own page count."""
    opts = Options(request_id="T", jobs=3)
    bodies = {f"log{n}.txt": "".join(f"{n} line {i}\n" for i in range(59 * (n + 1))).encode() for n in range(4)}
    with Converter(tmp_path / "w", opts) as conv:
        results = conv.convert_all(Attachment(name, body, "text/plain") for name, body in bodies.items())
        assert conv._render_pool is not None
    assert [(r.status, r.method, r.page_count) for r in results] == [("converted", "text", n + 1) for n in range(4)]
    assert "3 line 235" in PdfReader(str(results[3].pdf_path)).pages[3].extract_text()


def test_jobs_1_lays_out_inline(tmp_path):
    with Converter(tmp_path / "w", Options(request_id="T", jobs=1)) as conv:
        res = conv.convert(Attachment("a.txt", b"hello", "text/plain"))
        assert conv._render_pool is None
    assert (res.status, res.page_count) == ("converted", 1)


# ---------------------------------------------------------------------------
# Office names on bytes that are something else
# ---------------------------------------------------------------------------

#: What a reporting system exports as "Contact export.xls": UTF-16
#: tab-separated text. LibreOffice printed each row as one unbroken run.
TSV_AS_XLS = (
    "EmpID\tLastname\tFirstname\tPhonenumber\tJobTitle\tEmailAddress\r\n"
    "10000001\tDOE\tJANE\t5555550142\tAnalyst, Budget\tJANE.DOE@EXAMPLE.ORG\r\n"
    "10000002\tROE\tRICHARD\t5555550199\tAnalyst\t\r\n"
).encode("utf-16")


@pytest.mark.parametrize(
    "filename, ctype, content, expected",
    [
        ("Contact export.xls", "", TSV_AS_XLS, ("text", "txt")),
        ("export.xls", "", b"id,name\r\n1,Jane Doe\r\n", ("text", "txt")),
        ("report.doc", "", b"plain words in a .doc\n", ("text", "txt")),
        ("table.xls", "", b"<html><body><table><tr><td>Jane</td></tr></table></body></html>", ("html", "html")),
        ("table.xls", "", b"  <TABLE><TR><TD>Jane</TD></TR></TABLE>", ("html", "html")),
        ("renamed.xls", "", "XLSX", ("office", "xlsx")),
        ("renamed.docx", "", b"%PDF-1.7\n", ("pdf", "pdf")),
        ("real.xls", "", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, ("office", "xls")),
        ("real.xlsx", "", "XLSX", ("office", "xlsx")),
        ("biff.xls", "", b"\x09\x08\x10\x00\x00\x06\x05\x00" + b"\x00" * 64, ("office", "xls")),
        ("notes.rtf", "", b"not rtf at all", ("office", "rtf")),
        ("blob", "application/vnd.ms-excel", TSV_AS_XLS, ("text", "txt")),
    ],
)
def test_office_names_are_checked_against_the_bytes(filename, ctype, content, expected):
    if content == "XLSX":
        content = minimal_xlsx()
    assert detect(filename, ctype, content) == expected


def test_tab_separated_text_named_xls_prints_every_field_as_a_word(converter, stub_render_pdf):
    res = converter.convert(Attachment("Contact export.xls", TSV_AS_XLS, "application/vnd.ms-excel"))
    assert (res.status, res.method) == ("converted", "text")
    import pdfplumber

    with pdfplumber.open(res.pdf_path) as pdf:
        words = [w["text"] for w in pdf.pages[0].extract_words()]
    for word in ("DOE", "JANE", "ROE", "RICHARD", "Budget", "JANE.DOE@EXAMPLE.ORG"):
        assert word in words
