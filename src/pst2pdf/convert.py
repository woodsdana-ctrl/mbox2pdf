"""Attachment → PDF conversion (design contract: docs/DESIGN.md, "convert.py").

``Converter.convert`` turns one :class:`Attachment` into a
:class:`ConvertedAttachment` whose ``pdf_path`` is a sanitized PDF the
assembler can merge, or a placeholder carrying a reason. Results are cached
by sha256 so identical bytes are converted once per run, and the cache is
thread-safe so the CLI can push attachments through ``convert_all`` while
LibreOffice subprocesses run side by side. Page layout that would hold the
GIL (WeasyPrint for HTML and nested messages, reportlab for plain text) runs
in a pool of worker processes, so it uses every core too.

Dispatch is by file extension, then declared content type, then magic
bytes. Every path is wrapped so that any exception becomes a placeholder
rather than a crash: an archive with 3,000 attachments must never fail on
one of them.

Nothing here writes outside *work_dir*: LibreOffice gets a throwaway profile
and ``HOME`` under it, inputs are staged there, and outputs are moved into
``work_dir/converted``.
"""

from __future__ import annotations

import io
import logging
import multiprocessing
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import warnings
import zipfile
from collections.abc import Callable, Iterable
from concurrent.futures import Future, ProcessPoolExecutor, ThreadPoolExecutor
from datetime import datetime
from email.utils import parsedate_to_datetime
from html import escape
from pathlib import Path

from . import textpdf
from .model import Attachment, ConvertedAttachment, MessageModel, Options

try:  # POSIX only, and imported here so the fork hook never has to (see below)
    import resource
except ImportError:  # pragma: no cover - not a platform we ship
    resource = None  # type: ignore[assignment]

log = logging.getLogger(__name__)

__all__ = [
    "Converter",
    "detect",
    "sanitize_pdf",
    "pdf_page_count",
    "prepare_spreadsheet",
    "SANITIZE_EXCLUDED_KEYS",
    "PAGE_ALLOWED_KEYS",
]

# --------------------------------------------------------------------------
# PDF hygiene
# --------------------------------------------------------------------------

#: Keys dropped from every dictionary (page, XObject, font, ...) copied out of
#: an untrusted PDF. pypdf applies ``excluded_keys`` recursively during the
#: clone, so this reaches image XObjects and nested resources too.
#:
#: /AA page actions, /AF associated files (their EmbeddedFile streams would
#: otherwise survive), /Metadata XMP packets, /PieceInfo vendor private data,
#: /Thumb thumbnails, /B article beads, /StructParents structure-tree links,
#: /OC optional-content membership.
SANITIZE_EXCLUDED_KEYS: tuple[str, ...] = (
    "/AA", "/AF", "/Metadata", "/PieceInfo", "/Thumb", "/B", "/StructParents", "/OC",
)

#: The only keys a sanitized page dictionary may carry. Anything else,
#: including vendor keys such as Apple's ``/AAPL:PPK``, is dropped. ``/Annots``
#: is deliberately absent: widgets are flattened first, then every
#: annotation goes.
PAGE_ALLOWED_KEYS: frozenset[str] = frozenset(
    {
        "/Type", "/Parent", "/Contents", "/Resources", "/MediaBox", "/CropBox",
        "/BleedBox", "/TrimBox", "/ArtBox", "/Rotate", "/Group", "/UserUnit",
    }
)

# --------------------------------------------------------------------------
# Limits
# --------------------------------------------------------------------------

#: Virtual address space granted to each soffice process (Linux only).
SOFFICE_RLIMIT_AS = 4 * 1024**3
#: LibreOffice instances started at the same time, whatever ``options.jobs``
#: says; each one costs several hundred MB and a second or two of start-up.
SOFFICE_MAX_PARALLEL = 4
#: Lay out text, HTML and nested messages in worker processes when
#: ``options.jobs`` > 1. Tests that stub ``render.render_pdf`` in this process
#: turn it off, since a worker process would not see the stub.
RENDER_IN_PROCESSES = True
#: Text-like attachments larger than this go to a placeholder instead of a
#: multi-thousand-page ``<pre>`` render.
TEXT_MAX_BYTES = 5 * 1024 * 1024
#: Animated images stop producing pages after this many frames.
IMAGE_MAX_FRAMES = 200
#: Entries shown on a zip placeholder before "... and N more".
ZIP_LIST_LIMIT = 100

_LETTER = (612.0, 792.0)
_IMAGE_MARGIN = 36.0  # points

# --------------------------------------------------------------------------
# Dispatch tables
# --------------------------------------------------------------------------

_OFFICE_EXTS = {
    "doc", "docx", "docm", "dot", "dotx", "dotm",
    "xls", "xlsx", "xlsm", "xlsb", "xlt", "xltx",
    "ppt", "pptx", "pptm", "pps", "ppsx",
    "odt", "ods", "odp", "rtf", "wpd",
}
_IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp", "heic", "heif"}
_TEXT_EXTS = {"txt", "csv", "tsv", "log", "md", "json", "xml"}
_HTML_EXTS = {"html", "htm"}

_KIND_BY_EXT: dict[str, str] = {}
_KIND_BY_EXT.update({e: "office" for e in _OFFICE_EXTS})
_KIND_BY_EXT.update({e: "image" for e in _IMAGE_EXTS})
_KIND_BY_EXT.update({e: "text" for e in _TEXT_EXTS})
_KIND_BY_EXT.update({e: "html" for e in _HTML_EXTS})
_KIND_BY_EXT.update({"pdf": "pdf", "eml": "eml", "msg": "msg", "ics": "ics", "zip": "zip"})

_KIND_BY_CONTENT_TYPE: dict[str, tuple[str, str]] = {
    "application/pdf": ("pdf", "pdf"),
    "application/msword": ("office", "doc"),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ("office", "docx"),
    "application/vnd.ms-word.document.macroenabled.12": ("office", "docm"),
    "application/vnd.ms-excel": ("office", "xls"),
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ("office", "xlsx"),
    "application/vnd.ms-excel.sheet.macroenabled.12": ("office", "xlsm"),
    "application/vnd.ms-powerpoint": ("office", "ppt"),
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ("office", "pptx"),
    "application/vnd.oasis.opendocument.text": ("office", "odt"),
    "application/vnd.oasis.opendocument.spreadsheet": ("office", "ods"),
    "application/vnd.oasis.opendocument.presentation": ("office", "odp"),
    "application/rtf": ("office", "rtf"),
    "text/rtf": ("office", "rtf"),
    "application/wordperfect": ("office", "wpd"),
    "application/vnd.ms-outlook": ("msg", "msg"),
    "message/rfc822": ("eml", "eml"),
    "text/calendar": ("ics", "ics"),
    "text/html": ("html", "html"),
    "text/plain": ("text", "txt"),
    "text/csv": ("text", "csv"),
    "text/markdown": ("text", "md"),
    "text/xml": ("text", "xml"),
    "application/xml": ("text", "xml"),
    "application/json": ("text", "json"),
    "application/zip": ("zip", "zip"),
    "application/x-zip-compressed": ("zip", "zip"),
    "image/png": ("image", "png"),
    "image/jpeg": ("image", "jpg"),
    "image/gif": ("image", "gif"),
    "image/bmp": ("image", "bmp"),
    "image/tiff": ("image", "tiff"),
    "image/webp": ("image", "webp"),
}

#: Office names whose genuine files always come in a known container: OLE2
#: for the legacy binary formats, a zip for Office Open XML. Reports exported
#: from other systems are often tab-separated text, CSV or an HTML table saved
#: under one of these names. LibreOffice opens such a file as a spreadsheet
#: without splitting it at the tabs, so every row prints as one run of
#: characters ("ROERICHARD5555550199...") that a reviewer cannot read and
#: whole-word search cannot find a name in.
_OLE_EXTS = {"doc", "dot", "xls", "xlt", "ppt", "pps"}
_OOXML_EXTS = {"docx", "docm", "dotx", "dotm", "xlsx", "xlsm", "xltx", "pptx", "pptm", "ppsx"}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_ODF_MIMETYPES = {
    b"application/vnd.oasis.opendocument.text": "odt",
    b"application/vnd.oasis.opendocument.spreadsheet": "ods",
    b"application/vnd.oasis.opendocument.presentation": "odp",
}


def _kind_from_magic(content: bytes) -> tuple[str, str] | None:
    """Sniff the container format from leading bytes; None when unknown."""
    head = content[:16]
    if head.startswith(b"%PDF"):
        return "pdf", "pdf"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image", "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image", "jpg"
    if head.startswith((b"GIF87a", b"GIF89a")):
        return "image", "gif"
    if head.startswith(b"BM"):
        return "image", "bmp"
    if head.startswith((b"II*\x00", b"MM\x00*")):
        return "image", "tiff"
    if head.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image", "webp"
    if head.startswith(b"{\\rtf"):
        return "office", "rtf"
    if head.startswith(_OLE_MAGIC):
        return _kind_from_ole(content)
    if head.startswith(b"PK\x03\x04"):
        return _kind_from_zip(content)
    stripped = content[:64].lstrip()
    if stripped.startswith(b"BEGIN:VCALENDAR"):
        return "ics", "ics"
    if stripped[:15].lower().startswith((b"<!doctype html", b"<html")):
        return "html", "html"
    return None


def _kind_from_ole(content: bytes) -> tuple[str, str] | None:
    try:
        import olefile
    except ImportError:  # pragma: no cover - olefile ships with extract-msg
        return None
    try:
        with olefile.OleFileIO(io.BytesIO(content)) as ole:
            if ole.exists("__properties_version1.0"):
                return "msg", "msg"
            if ole.exists("WordDocument"):
                return "office", "doc"
            if ole.exists("Workbook") or ole.exists("Book"):
                return "office", "xls"
            if ole.exists("PowerPoint Document"):
                return "office", "ppt"
    except Exception:
        return None
    return None


def _kind_from_zip(content: bytes) -> tuple[str, str] | None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            names = set(zf.namelist())
            if "[Content_Types].xml" in names:
                if any(n.startswith("word/") for n in names):
                    return "office", "docx"
                if any(n.startswith("xl/") for n in names):
                    return "office", "xlsx"
                if any(n.startswith("ppt/") for n in names):
                    return "office", "pptx"
            if "mimetype" in names:
                ext = _ODF_MIMETYPES.get(zf.read("mimetype").strip())
                if ext:
                    return "office", ext
    except Exception:
        return None
    return "zip", "zip"


def _sniff_text(content: bytes) -> str | None:
    """The start of *content* as text when it reads as text (a byte-order
    mark, or no NUL bytes and almost nothing unprintable); else None."""
    sample = content[:65536]
    if not sample.startswith((b"\xff\xfe", b"\xfe\xff")) and b"\x00" in sample:
        return None
    text = _decode_text(sample)
    if not text or "\x00" in text:
        return None
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\t\r\n\f")
    return text if printable >= 0.97 * len(text) else None


def _verify_office(ext: str, content: bytes) -> tuple[str, str]:
    """``("office", ext)`` unless the bytes are plainly not that format: then
    whatever they really are (another Office format, a PDF, HTML, or text).
    Bytes nothing recognizes still go to LibreOffice, which reads old formats
    (BIFF, SYLK) this does not."""
    if ext not in _OLE_EXTS | _OOXML_EXTS:
        return "office", ext
    if content.startswith(_OLE_MAGIC if ext in _OLE_EXTS else b"PK\x03\x04"):
        return "office", ext
    hit = _kind_from_magic(content)
    if hit:
        return hit
    text = _sniff_text(content)
    if text is not None:
        if text.lstrip()[:15].lower().startswith(("<!doctype html", "<html", "<table")):
            return "html", "html"
        return "text", "txt"
    return "office", ext


def detect(filename: str, content_type: str, content: bytes) -> tuple[str | None, str]:
    """Return ``(kind, extension)`` for an attachment.

    *kind* is one of ``office image pdf text html eml msg ics zip`` or None
    when nothing recognizes the file. *extension* is the lower-case extension
    (without dot) the conversion should treat the bytes as; for a recognized
    kind it is always a member of the dispatch table, so it is safe to use as
    the staging filename's suffix.
    """
    ext = Path(filename or "").suffix.lower().lstrip(".")
    kind = _KIND_BY_EXT.get(ext)
    if kind == "office":
        return _verify_office(ext, content)
    if kind:
        return kind, ext
    ctype = (content_type or "").split(";", 1)[0].strip().lower()
    hit = _KIND_BY_CONTENT_TYPE.get(ctype)
    if hit and hit[0] == "office":
        return _verify_office(hit[1], content)
    if hit:
        return hit
    hit = _kind_from_magic(content)
    if hit:
        return hit
    return None, ext


# --------------------------------------------------------------------------
# PDF helpers (pypdf)
# --------------------------------------------------------------------------


class _Placeholder(Exception):
    """Raised inside a conversion to abandon it with a user-facing reason."""


def pdf_page_count(path: Path) -> int:
    from pypdf import PdfReader

    return len(PdfReader(str(path)).pages)


def _flatten_widgets(writer, page) -> int:
    """Draw each visible form widget's normal appearance into the page content.

    Follows the PDF 1.7 appearance-stream algorithm (12.5.5): the form
    XObject's BBox, transformed by its Matrix, is fitted to the annotation
    Rect. A synthetic page carrying the appearance stream is merged in with
    pypdf's public page-merge API, so no private writer internals are used.
    Returns how many widgets were flattened. Widgets without an appearance
    stream (or hidden ones) are simply dropped by the caller.
    """
    from pypdf import PageObject, Transformation
    from pypdf.generic import ArrayObject, ContentStream, DictionaryObject, NameObject

    annots = page.get("/Annots")
    if not annots:
        return 0
    flattened = 0
    for ref in list(annots):
        try:
            annot = ref.get_object()
            if annot.get("/Subtype") != "/Widget":
                continue
            flags = int(annot.get("/F", 0))
            if flags & 2 or flags & 32:  # Hidden, NoView
                continue
            ap = annot.get("/AP")
            if not ap:
                continue
            normal = ap.get_object().get("/N")
            if normal is None:
                continue
            normal = normal.get_object()
            if not hasattr(normal, "get_data"):  # appearance sub-dictionary keyed by state
                state = annot.get("/AS")
                if state is None or state not in normal:
                    continue
                normal = normal[state].get_object()
            rect = [float(v) for v in annot["/Rect"]]
            bbox = [float(v) for v in normal.get("/BBox", [0, 0, 0, 0])]
            matrix = [float(v) for v in normal.get("/Matrix", [1, 0, 0, 1, 0, 0])]
            # Transform the bbox corners by Matrix and take their bounds.
            xs, ys = [], []
            for x, y in ((bbox[0], bbox[1]), (bbox[2], bbox[1]), (bbox[2], bbox[3]), (bbox[0], bbox[3])):
                xs.append(matrix[0] * x + matrix[2] * y + matrix[4])
                ys.append(matrix[1] * x + matrix[3] * y + matrix[5])
            tw, th = max(xs) - min(xs), max(ys) - min(ys)
            rw, rh = abs(rect[2] - rect[0]), abs(rect[3] - rect[1])
            sx = rw / tw if tw else 1.0
            sy = rh / th if th else 1.0
            a = Transformation((sx, 0, 0, sy, min(rect[0], rect[2]) - min(xs) * sx, min(rect[1], rect[3]) - min(ys) * sy))
            ctm = Transformation(tuple(matrix)).transform(a)

            synth = PageObject.create_blank_page(writer, max(tw, 1.0), max(th, 1.0))
            synth[NameObject("/MediaBox")] = ArrayObject([*map(_num, bbox)])
            resources = normal.get("/Resources")
            if resources is not None:
                synth[NameObject("/Resources")] = resources.get_object()
            else:
                synth[NameObject("/Resources")] = DictionaryObject()
            synth.replace_contents(ContentStream(normal, writer))
            page.merge_transformed_page(synth, ctm, over=True)
            flattened += 1
        except Exception as exc:  # one bad widget must not sink the page
            log.debug("widget flatten skipped: %s", exc)
    return flattened


def _num(value: float):
    from pypdf.generic import FloatObject, NumberObject

    return NumberObject(int(value)) if float(value).is_integer() else FloatObject(value)


def sanitize_pdf(content: bytes, out_path: Path) -> dict[str, int]:
    """Rebuild *content* page by page at *out_path* without active content.

    Only page objects are copied into a fresh writer, so document-level
    JavaScript, OpenAction, the embedded-files name tree and the AcroForm
    never make it across. Page-level channels go too: the keys in
    :data:`SANITIZE_EXCLUDED_KEYS` are excluded recursively (so associated
    files, XMP on pages and on image XObjects, PieceInfo and thumbnails never
    reach the writer), form widgets are flattened where they carry an
    appearance stream, every remaining annotation (links, notes, file
    attachments, widgets) is removed, and finally each page dictionary is
    reduced to :data:`PAGE_ALLOWED_KEYS` so unknown vendor keys cannot pass.
    Raises :class:`_Placeholder` for a PDF that needs a password.
    """
    from pypdf import PdfReader, PdfWriter
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
    except PdfReadError as exc:
        raise _Placeholder(f"not a readable PDF ({exc})") from exc
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")
        except Exception:  # missing cipher support counts as locked
            opened = 0
        if not opened:
            raise _Placeholder("password protected")

    stats = {"pages": 0, "flattened": 0, "annotations": 0}
    writer = PdfWriter()
    for src in reader.pages:
        page = writer.add_page(src, excluded_keys=SANITIZE_EXCLUDED_KEYS)
        annots = page.get("/Annots")
        if annots is not None:
            annots = annots.get_object()
        if annots:
            stats["annotations"] += len(annots)
            stats["flattened"] += _flatten_widgets(writer, page)
        stats["pages"] += 1
    writer.remove_annotations(subtypes=None)
    for page in writer.pages:
        for key in list(page.keys()):
            if key not in PAGE_ALLOWED_KEYS:
                del page[key]
    writer.add_metadata({"/Producer": "pst2pdf"})
    with out_path.open("wb") as fh:
        writer.write(fh)
    return stats


# --------------------------------------------------------------------------
# Text / HTML wrappers
# --------------------------------------------------------------------------

_PAGE_CSS = """
@page { size: Letter; margin: 0.75in; }
body { font-family: "DejaVu Sans", "Noto Sans CJK SC", sans-serif; font-size: 10pt; line-height: 1.35;
       font-variant-ligatures: none; hyphens: none; overflow-wrap: break-word; }
pre { font-family: "DejaVu Sans Mono", "Noto Sans Mono CJK SC", monospace; font-size: 8.5pt;
      white-space: pre-wrap; overflow-wrap: anywhere; margin: 0; }
img { max-width: 100% !important; max-height: 9in !important; }
table { border-collapse: collapse; }
td, th { vertical-align: top; }
a { color: inherit; text-decoration: none; }
"""


def _wrap_html(fragment: str) -> str:
    return (
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\">"
        f"<style>{_PAGE_CSS}</style></head><body>{fragment}</body></html>"
    )


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", "replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", "replace")


def _html_fragment(document: str) -> str:
    """Reduce a full HTML document to its body content (or the whole text)."""
    import re

    match = re.search(r"<body[^>]*>(.*)</body>", document, re.IGNORECASE | re.DOTALL)
    if match:
        return match.group(1)
    return re.sub(
        r"<!DOCTYPE[^>]*>|</?html[^>]*>|<head[^>]*>.*?</head>",
        "",
        document,
        flags=re.IGNORECASE | re.DOTALL,
    )


def _clean_html(fragment: str) -> str | None:
    """Run parse.clean_html if it exists; None when the helper is unavailable."""
    try:
        from .parse import clean_html
    except ImportError:
        return None
    return clean_html(fragment)


# --------------------------------------------------------------------------
# Converter
# --------------------------------------------------------------------------


class Converter:
    """Convert attachments to PDF with a per-run, sha256-keyed cache.

    Nested messages (``.eml``/``.msg``/``.ics``) are rendered through the same
    template as top-level messages, in the worker pool. A ``render_message``
    callback, when given, renders them instead (inline, one at a time) and
    returns the PDF path; the Converter moves that file into its cache
    directory.
    """

    def __init__(
        self,
        work_dir: Path,
        options: Options,
        render_message: Callable[[MessageModel], Path] | None = None,
    ) -> None:
        self.work_dir = Path(work_dir)
        self.options = options
        self.render_message = render_message
        self.out_dir = self.work_dir / "converted"
        self.tmp_root = self.work_dir / "convert-tmp"
        self.out_dir.mkdir(parents=True, exist_ok=True)
        self.tmp_root.mkdir(parents=True, exist_ok=True)

        jobs = options.jobs or os.cpu_count() or 1
        self.jobs = max(1, int(jobs))
        self._pool = ThreadPoolExecutor(max_workers=self.jobs, thread_name_prefix="convert")
        self._soffice_slots = threading.BoundedSemaphore(min(self.jobs, SOFFICE_MAX_PARALLEL))
        # WeasyPrint (Pango/Cairo) is not documented as thread-safe, and the
        # supplied render_message may use it inline, so in-process layout
        # serializes. With jobs > 1 layout runs in worker processes instead.
        self._render_lock = threading.Lock()
        self._render_pool: ProcessPoolExecutor | None = None
        self._render_pool_lock = threading.Lock()
        self._cache_lock = threading.Lock()
        self._cache: dict[str, Future[ConvertedAttachment]] = {}
        self.blocked_urls = 0  # remote resources refused while laying out HTML attachments
        self.soffice = shutil.which("soffice") or shutil.which("libreoffice")

    # -- public API -------------------------------------------------------

    def convert(self, attachment: Attachment) -> ConvertedAttachment:
        """Convert (or fetch from cache) synchronously. Safe from any thread."""
        return self.submit(attachment).result()

    def submit(self, attachment: Attachment) -> Future[ConvertedAttachment]:
        """Schedule a conversion on the pool; identical bytes share one Future."""
        digest = attachment.sha256
        with self._cache_lock:
            fut = self._cache.get(digest)
            if fut is None:
                fut = self._pool.submit(self._convert_uncached, attachment, digest)
                self._cache[digest] = fut
        return fut

    def convert_all(self, attachments: Iterable[Attachment]) -> list[ConvertedAttachment]:
        """Convert many attachments concurrently; results keep input order."""
        futures = [self.submit(a) for a in attachments]
        return [f.result() for f in futures]

    def close(self) -> None:
        self._pool.shutdown(wait=True)
        if self._render_pool is not None:
            self._render_pool.shutdown(wait=True)
        shutil.rmtree(self.tmp_root, ignore_errors=True)

    # -- layout in worker processes -------------------------------------------

    def _pooled(self) -> bool:
        return RENDER_IN_PROCESSES and self.jobs > 1

    def _layout(self, fn: Callable, *args):
        """Run the module-level layout function *fn* in the worker pool, or
        inline (one at a time) with ``--jobs 1``. Called from converter threads."""
        if not self._pooled():
            with self._render_lock:
                return fn(*args)
        with self._render_pool_lock:
            if self._render_pool is None:
                # Not fork: the parent is multi-threaded by now (converter
                # threads, LibreOffice waits), and a forked child can inherit
                # a lock some other thread held. forkserver forks from a clean
                # single-threaded server.
                method = "forkserver" if "forkserver" in multiprocessing.get_all_start_methods() else "spawn"
                self._render_pool = ProcessPoolExecutor(
                    max_workers=self.jobs, mp_context=multiprocessing.get_context(method)
                )
        return self._render_pool.submit(fn, *args).result()

    def _weasyprint(self, html: str, out: Path, name: str = "") -> Path:
        from . import render

        result = self._layout(render.render_pdf, html, out)
        with self._cache_lock:
            self.blocked_urls += result.blocked
            render.note_offpage(result, name or "attachment")
        return out

    def __enter__(self) -> Converter:
        return self

    def __exit__(self, *exc) -> None:
        self.close()

    # -- dispatch ---------------------------------------------------------

    def _convert_uncached(self, attachment: Attachment, digest: str) -> ConvertedAttachment:
        name = attachment.filename or "attachment"
        try:
            if attachment.size == 0:
                raise _Placeholder("empty file (0 bytes)")
            if attachment.size > self.options.attachment_max_bytes:
                raise _Placeholder(
                    f"file is {_mb(attachment.size)} MB, over the "
                    f"{_mb(self.options.attachment_max_bytes)} MB conversion limit"
                )
            kind, ext = detect(name, attachment.content_type, attachment.content)
            named = Path(name).suffix.lower().lstrip(".")
            if named in _OLE_EXTS | _OOXML_EXTS and (kind, ext) != ("office", named):
                log.info("%s: named .%s but the bytes are %s; converting it as that", name, named, ext)
            if kind is None:
                raise _Placeholder(
                    f"unsupported file type ({'.' + ext if ext else attachment.content_type or 'unknown'})"
                )
            if kind == "zip":
                raise _Placeholder(self._zip_listing(attachment.content, expanded=self.options.expand_zips and "/" not in attachment.filename))
            method, produce = {
                "office": ("libreoffice", self._via_soffice),
                "image": ("image", self._via_image),
                "pdf": ("pdf", self._via_pdf),
                "text": ("text", self._via_text),
                "html": ("html", self._via_html),
                "eml": ("message", self._via_eml),
                "msg": ("message", self._via_msg),
                "ics": ("calendar", self._via_ics),
            }[kind]
            final = self.out_dir / f"{digest}.pdf"
            tmp = Path(tempfile.mkdtemp(prefix="cv-", dir=self.tmp_root))
            try:
                produced, method = produce(attachment, ext, tmp, method)
                if produced != final:
                    shutil.move(str(produced), str(final))
            finally:
                shutil.rmtree(tmp, ignore_errors=True)
            pages = pdf_page_count(final)
            if pages == 0:
                final.unlink(missing_ok=True)
                raise _Placeholder("conversion produced no pages")
            max_pages = self.options.attachment_max_pages
            if max_pages and pages > max_pages:
                final.unlink(missing_ok=True)
                raise _Placeholder(f"{pages:,} pages, more than the {max_pages:,}-page limit")
            status = "passthrough" if method == "pdf" else "converted"
            log.debug("%s: %s via %s, %d page(s)", name, status, method, pages)
            return ConvertedAttachment(attachment, status, method, final, pages)
        except _Placeholder as exc:
            log.info("%s: placeholder: %s", name, str(exc).splitlines()[0])
            return ConvertedAttachment(attachment, "placeholder", "none", None, 0, str(exc))
        except Exception as exc:  # noqa: BLE001 - anything else becomes a placeholder
            log.warning("%s: conversion failed: %s: %s", name, type(exc).__name__, exc)
            reason = f"conversion failed ({type(exc).__name__}: {str(exc)[:200]})"
            return ConvertedAttachment(attachment, "placeholder", "none", None, 0, reason)

    # -- LibreOffice --------------------------------------------------------

    def _via_soffice(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        if not self.soffice:
            raise _Placeholder("LibreOffice not available")
        content = att.content
        if ext in _FIT_SPREADSHEET_EXTS:
            fitted = prepare_spreadsheet(content, att.filename)
            if fitted is not None:
                content, ext = fitted, "xlsx"
            else:
                log.info("%s: could not pre-process the workbook, converting it as is", att.filename)
        src = tmp / f"input.{ext}"
        src.write_bytes(content)
        profile = tmp / "profile"
        _write_profile(profile)
        out_dir = tmp / "out"
        out_dir.mkdir()
        cmd = self._soffice_command(src, out_dir, profile)
        env = {
            "HOME": str(tmp),
            "PATH": os.environ.get("PATH", "/usr/bin:/bin"),
            "LANG": os.environ.get("LANG", "C.UTF-8"),
            "SAL_USE_VCLPLUGIN": "headless",
            "TMPDIR": str(tmp),
        }
        with self._soffice_slots:
            result = _run_killable(cmd, cwd=tmp, env=env, timeout=self.options.attachment_timeout)
        if result is None:
            raise _Placeholder(f"conversion timed out after {self.options.attachment_timeout} s")
        code, stderr = result
        produced = out_dir / "input.pdf"
        if code != 0 or not produced.is_file():
            tail = stderr.strip().splitlines()[-1] if stderr.strip() else f"exit status {code}"
            raise _Placeholder(f"LibreOffice could not convert this file ({tail[:160]})")
        # Office exports carry link annotations (and can carry attached files);
        # the same pass-through hygiene as for PDF attachments applies.
        clean = tmp / "clean.pdf"
        sanitize_pdf(produced.read_bytes(), clean)
        return clean, method

    def _soffice_command(self, src: Path, out_dir: Path, profile: Path) -> list[str]:
        return [
            self.soffice,
            "--headless",
            "--invisible",
            "--norestore",
            "--nolockcheck",
            "--nologo",
            "--nodefault",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(src),
        ]

    # -- images -------------------------------------------------------------

    def _via_image(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        from PIL import Image, ImageSequence
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas

        from . import images  # registers the HEIF opener when available

        out = tmp / "image.pdf"
        try:
            opened = Image.open(io.BytesIO(att.content))
        except Image.DecompressionBombError as exc:
            raise _Placeholder(f"image is too large to open safely ({exc})") from exc
        with opened as im:
            w0, h0 = im.size
            limit = Image.MAX_IMAGE_PIXELS
            if limit and w0 * h0 > limit:
                raise _Placeholder(
                    f"image is {w0}x{h0} px ({w0 * h0:,} pixels), over the {limit:,}-pixel limit"
                )
            pdf = canvas.Canvas(str(out))
            pdf.setAuthor("")
            pdf.setTitle("")
            pdf.setSubject("")
            pdf.setCreator("pst2pdf")
            frames = 0
            for frame in ImageSequence.Iterator(im):
                if frames >= IMAGE_MAX_FRAMES:
                    log.info("%s: only the first %d frames rendered", att.filename, IMAGE_MAX_FRAMES)
                    break
                rgb = _to_rgb(frame)
                if max(rgb.size) > images.MAX_SIDE:
                    scale = images.MAX_SIDE / max(rgb.size)
                    rgb = rgb.resize((max(1, round(rgb.width * scale)), max(1, round(rgb.height * scale))), Image.LANCZOS)
                w, h = rgb.size
                page_w, page_h = _LETTER if h >= w else (_LETTER[1], _LETTER[0])
                pdf.setPageSize((page_w, page_h))
                box_w, box_h = page_w - 2 * _IMAGE_MARGIN, page_h - 2 * _IMAGE_MARGIN
                scale = min(box_w / w, box_h / h)
                draw_w, draw_h = w * scale, h * scale
                x = (page_w - draw_w) / 2
                y = (page_h - draw_h) / 2
                if w * h > images.JPEG_PIXELS:
                    # A photo stored as raw pixels is ~40 MB a page; hand
                    # reportlab a JPEG file so it embeds the compressed stream.
                    jpg = tmp / f"frame-{frames}.jpg"
                    jpg.write_bytes(images.encode_jpeg(rgb))
                    pdf.drawImage(str(jpg), x, y, width=draw_w, height=draw_h)
                else:
                    pdf.drawImage(ImageReader(rgb), x, y, width=draw_w, height=draw_h)
                pdf.showPage()
                frames += 1
            if frames == 0:
                raise _Placeholder("image has no frames")
            pdf.save()
        return out, method

    # -- PDF ----------------------------------------------------------------

    def _via_pdf(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        out = tmp / "clean.pdf"
        stats = sanitize_pdf(att.content, out)
        if stats["annotations"]:
            log.info(
                "%s: removed %d annotation(s), flattened %d form widget(s)",
                att.filename, stats["annotations"], stats["flattened"],
            )
        return out, method

    # -- text / HTML ------------------------------------------------------

    def _via_text(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        if att.size > TEXT_MAX_BYTES:
            raise _Placeholder(
                f"text file is {_mb(att.size)} MB, over the {_mb(TEXT_MAX_BYTES)} MB render limit"
            )
        raw = _decode_text(att.content)
        text = textpdf.normalize(raw)
        if textpdf.covers(text):
            out = tmp / "text.pdf"
            # Delimited exports print as an aligned table. Commas count only
            # for a .csv or an export under a spreadsheet's name: prose has them.
            named = Path(att.filename or "").suffix.lower().lstrip(".")
            delimiter = textpdf.table_delimiter(raw, allow_comma=named == "csv" or named in _OLE_EXTS | _OOXML_EXTS)
            if delimiter:
                self._layout(textpdf.render_table_pdf, raw, delimiter, out)
            else:
                self._layout(textpdf.render_text_pdf, text, out)
            return out, method
        # Characters the monospace font cannot draw: WeasyPrint's font
        # fallback can, so nothing is left as an empty box.
        return self._render_html(f"<pre>{escape(text)}</pre>", tmp, att.filename), method

    def _via_html(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        if att.size > TEXT_MAX_BYTES:
            raise _Placeholder(
                f"HTML file is {_mb(att.size)} MB, over the {_mb(TEXT_MAX_BYTES)} MB render limit"
            )
        source = _decode_text(att.content)
        cleaned = _clean_html(_html_fragment(source))
        if cleaned is None:
            log.warning("%s: parse.clean_html unavailable, rendering HTML source", att.filename)
            fragment = f"<pre>{escape(source)}</pre>"
        else:
            fragment = cleaned
        return self._render_html(fragment, tmp, att.filename), method

    def _render_html(self, fragment: str, tmp: Path, name: str = "") -> Path:
        return self._weasyprint(_wrap_html(fragment), tmp / "text.pdf", name)

    # -- nested messages ----------------------------------------------------

    def _via_eml(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        from .parse import parse_eml

        src = tmp / "nested.eml"
        src.write_bytes(att.content)
        model = parse_eml(src, readpst_compat=self.options.input_format != "mbox")
        if att.filename and not model.subject:
            model.subject = Path(att.filename).stem
        return self._render_model(model, tmp, att.filename), method

    def _via_msg(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        src = tmp / "nested.msg"
        src.write_bytes(att.content)
        model = _msg_to_model(src, att.filename)
        return self._render_model(model, tmp, att.filename), method

    def _via_ics(self, att: Attachment, ext: str, tmp: Path, method: str) -> tuple[Path, str]:
        src = tmp / "nested.ics"
        src.write_bytes(att.content)
        try:
            from .calendar import parse_ics
        except ImportError:
            log.warning("%s: calendar module unavailable, rendering ICS as text", att.filename)
            path, _ = self._via_text(att, ext, tmp, "text")
            return path, "text"
        model = parse_ics(src)
        return self._render_model(model, tmp, att.filename), method

    def _render_model(self, model: MessageModel, tmp: Path, name: str = "") -> Path:
        if self.options.display_timezone and model.date and model.date.tzinfo:
            from zoneinfo import ZoneInfo
            model.date = model.date.astimezone(ZoneInfo(self.options.display_timezone))
        if self.render_message is None:
            from . import render

            return self._weasyprint(render.build_html(model), tmp / "message.pdf", name)
        with self._render_lock:
            produced = Path(self.render_message(model))
        if not produced.is_file():
            raise RuntimeError(f"render_message returned no file: {produced}")
        # Take ownership of the PDF so the cache directory holds every result.
        target = tmp / "message.pdf"
        shutil.move(str(produced), str(target))
        return target

    # -- zip ----------------------------------------------------------------

    @staticmethod
    def _zip_listing(content: bytes, expanded: bool = False) -> str:
        from .render import human_size

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                infos = zf.infolist()
        except zipfile.BadZipFile as exc:
            raise _Placeholder(f"not a readable zip archive ({exc})") from exc
        total = sum(i.file_size for i in infos)
        if expanded:
            lines = [f"it is a zip archive of {len(infos)} entries ({human_size(total)}); the files inside follow as their own attachments, except any noted here as password protected"]
        else:
            lines = [f"zip archive not expanded: {len(infos)} entries, {human_size(total)} uncompressed"]
        for info in infos[:ZIP_LIST_LIMIT]:
            flag = " (encrypted)" if info.flag_bits & 0x1 else ""
            if info.is_dir():
                lines.append(f"{info.filename}")
            else:
                lines.append(f"{info.filename} ({human_size(info.file_size)}){flag}")
        if len(infos) > ZIP_LIST_LIMIT:
            lines.append(f"... and {len(infos) - ZIP_LIST_LIMIT} more")
        return "\n".join(lines)


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------


def _mb(num: int) -> str:
    return f"{num / (1024 * 1024):.1f}".rstrip("0").rstrip(".")


#: Workbook formats openpyxl can rewrite before LibreOffice sees them. Legacy
#: ``.xls`` is binary and is left alone.
_FIT_SPREADSHEET_EXTS = frozenset({"xlsx", "xlsm"})
#: ``warnings.catch_warnings`` swaps process-global state and attachments
#: convert on a thread pool, so only one pre-pass may hold it at a time.
#: openpyxl is pure Python and keeps the GIL throughout, so serializing the
#: rewrite costs next to nothing.
_SPREADSHEET_WARNING_LOCK = threading.Lock()
#: A sheet wider than this many used columns is printed landscape.
SPREADSHEET_LANDSCAPE_COLUMNS = 8
_PAPER_LETTER = 1  # ECMA-376 paperSize code


def prepare_spreadsheet(content: bytes, name: str = "") -> bytes | None:
    """Rewrite an ``.xlsx``/``.xlsm`` so LibreOffice prints it readably.

    LibreOffice honours whatever page setup the workbook stores, which for a
    file nobody ever printed means one column-strip per page, rows cut off
    at stored heights, and thousands of pages for a modest sheet. Every
    worksheet is set to fit its width to one Letter page (as many pages tall
    as it takes), landscape when it is wider than
    :data:`SPREADSHEET_LANDSCAPE_COLUMNS` columns, with stored custom row
    heights cleared so wrapped cells grow to show their text. Returns the
    rewritten workbook (always ``.xlsx``; VBA is not carried over), or None
    when openpyxl cannot read or write the file, in which case the caller
    converts the original bytes.

    openpyxl warns on stderr about workbook parts it cannot round-trip (an
    unparsable printed header/footer, the data-validation extension). Those
    warnings are captured and logged against *name* at debug level instead,
    so they stay out of the wrapper's progress output, and so a run reports
    one line per affected workbook rather than Python's one line per
    warning site per process.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.worksheet.properties import PageSetupProperties
    except ImportError:  # pragma: no cover - openpyxl is a declared dependency
        return None
    label = name or "workbook"
    try:
        with _SPREADSHEET_WARNING_LOCK, warnings.catch_warnings(record=True) as caught:
            # "always", so a second workbook with the same defect is reported
            # too: the default filter shows each warning site once per process.
            warnings.simplefilter("always")
            wb = load_workbook(io.BytesIO(content))
            for ws in wb.worksheets:
                setup = ws.page_setup
                setup.fitToWidth = 1
                setup.fitToHeight = 0
                setup.paperSize = _PAPER_LETTER
                if ws.max_column > SPREADSHEET_LANDSCAPE_COLUMNS:
                    setup.orientation = "landscape"
                props = ws.sheet_properties
                if props.pageSetUpPr is None:
                    props.pageSetUpPr = PageSetupProperties(fitToPage=True)
                else:
                    props.pageSetUpPr.fitToPage = True
                for dim in ws.row_dimensions.values():
                    if dim.customHeight:
                        dim.height = None
            out = io.BytesIO()
            wb.save(out)
        for entry in caught:
            log.debug(
                "%s: openpyxl dropped part of the workbook: %s",
                label,
                " ".join(str(entry.message).split()),
            )
        return out.getvalue()
    except Exception as exc:  # noqa: BLE001 - any failure means "convert the original"
        log.debug("%s: spreadsheet pre-processing skipped: %s: %s", label, type(exc).__name__, exc)
        return None


def _to_rgb(frame):
    """Flatten transparency onto white and normalize the mode for reportlab."""
    from PIL import Image

    if frame.mode in ("RGBA", "LA") or (frame.mode == "P" and "transparency" in frame.info):
        rgba = frame.convert("RGBA")
        background = Image.new("RGB", rgba.size, (255, 255, 255))
        background.paste(rgba, mask=rgba.getchannel("A"))
        return background
    if frame.mode == "RGB":
        return frame.copy()
    return frame.convert("RGB")


_PROFILE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
           xmlns:xs="http://www.w3.org/2001/XMLSchema"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <item oor:path="/org.openoffice.Office.Common/Security/Scripting">
    <prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Security/Scripting">
    <prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Misc">
    <prop oor:name="FirstRun" oor:op="fuse"><value>false</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Misc">
    <prop oor:name="ShowTipOfTheDay" oor:op="fuse"><value>false</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Security/Scripting">
    <prop oor:name="HyperlinksWithCtrlClick" oor:op="fuse"><value>true</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Security/Scripting">
    <prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop>
  </item>
  <item oor:path="/org.openoffice.Office.Common/Save/Document">
    <prop oor:name="AutoSave" oor:op="fuse"><value>false</value></prop>
  </item>
  <!-- Writer: 0 = never update links, 1 = on request, 2 = always. -->
  <item oor:path="/org.openoffice.Office.Writer/Content/Update">
    <prop oor:name="Link" oor:op="fuse"><value>0</value></prop>
  </item>
  <!-- Calc (ScLkUpdMode): 0 = ALWAYS, 1 = NEVER, 2 = on demand. Not the same scale as Writer. -->
  <item oor:path="/org.openoffice.Office.Calc/Content/Update">
    <prop oor:name="Link" oor:op="fuse"><value>1</value></prop>
  </item>
</oor:items>
"""


def _write_profile(profile: Path) -> None:
    """Create a fresh LibreOffice user profile whose registry already carries
    MacroSecurityLevel=3 (Very high: only macros from trusted locations, of
    which there are none), macro execution disabled, untrusted referer links
    blocked, and external link updates off in both Writer and Calc, so a
    booby-trapped document opened headless cannot run code or phone home."""
    (profile / "user").mkdir(parents=True, exist_ok=True)
    (profile / "user" / "registrymodifications.xcu").write_text(_PROFILE_XCU, encoding="utf-8")


def _limit_address_space() -> None:  # pragma: no cover - runs in the child
    """Cap the child's address space, between fork and exec.

    The parent is multi-threaded (the conversion pool), and only the forking
    thread survives into the child: anything still holding a lock at fork time
    holds it forever there. So this does one bare syscall and nothing else --
    importing ``resource`` here would take the import lock and could hang the
    child until the conversion timeout killed it.
    """
    resource.setrlimit(resource.RLIMIT_AS, (SOFFICE_RLIMIT_AS, SOFFICE_RLIMIT_AS))


def _run_killable(cmd: list[str], cwd: Path, env: dict[str, str], timeout: int) -> tuple[int, str] | None:
    """Run *cmd* in its own session; on timeout kill the whole process group.

    soffice is a launcher (``oosplash``) that forks ``soffice.bin``, so killing
    only the direct child would leave the real worker running. Returns
    ``(returncode, stderr)`` or None when the timeout fired.
    """
    preexec = _limit_address_space if sys.platform.startswith("linux") and resource else None
    proc = subprocess.Popen(
        cmd,
        cwd=str(cwd),
        env=env,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        start_new_session=True,
        preexec_fn=preexec,
    )
    try:
        _, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        _kill_group(proc)
        return None
    return proc.returncode, err.decode("utf-8", "replace")


def _kill_group(proc: subprocess.Popen) -> None:
    try:
        os.killpg(proc.pid, signal.SIGKILL)
    except ProcessLookupError:
        pass
    except PermissionError:  # pragma: no cover
        proc.kill()
    try:
        proc.communicate(timeout=10)
    except subprocess.TimeoutExpired:  # pragma: no cover
        proc.kill()
        proc.communicate()


def _msg_to_model(path: Path, filename: str) -> MessageModel:
    """Load an Outlook .msg with extract_msg into a MessageModel.

    Inline images referenced by ``cid:`` become ``data:`` URIs as in
    parse.py; other attachments are listed (with their bytes, so the header
    shows sizes) but never converted, per the one-level rule.
    """
    import base64

    import extract_msg

    msg = extract_msg.openMsg(str(path))
    try:
        subject = _s(getattr(msg, "subject", "")) or Path(filename or "").stem
        html = getattr(msg, "htmlBody", None)
        if html:
            body = _html_fragment(_decode_text(html) if isinstance(html, bytes) else str(html))
            cleaned = _clean_html(body)
            body_html = cleaned if cleaned is not None else body
        else:
            plain = getattr(msg, "body", None)
            if isinstance(plain, bytes):
                plain = _decode_text(plain)
            body_html = f"<pre>{escape(plain)}</pre>" if plain else "<p><em>(no message body)</em></p>"

        attachments: list[Attachment] = []
        has_images = False
        for att in getattr(msg, "attachments", None) or []:
            data = getattr(att, "data", None)
            name = _s(getattr(att, "longFilename", None) or getattr(att, "name", None) or getattr(att, "shortFilename", None))
            if data is None:
                continue
            if not isinstance(data, (bytes, bytearray)):
                export = getattr(data, "exportBytes", None)
                if not callable(export):
                    continue
                data = export()
                name = name or "attachment"
                if not name.lower().endswith(".msg"):
                    name += ".msg"
            data = bytes(data)
            mimetype = _s(getattr(att, "mimetype", None)) or "application/octet-stream"
            cid = _s(getattr(att, "cid", None) or getattr(att, "contentId", None)).strip("<>")
            if cid and f"cid:{cid}" in body_html:
                data_uri = f"data:{mimetype};base64,{base64.b64encode(data).decode('ascii')}"
                body_html = body_html.replace(f"cid:{cid}", data_uri)
                has_images = True
                continue
            attachments.append(Attachment(name or "attachment", data, mimetype))

        raw_date = getattr(msg, "date", None)
        date: datetime | None
        if isinstance(raw_date, datetime):
            date = raw_date
        else:
            try:
                date = parsedate_to_datetime(str(raw_date)) if raw_date else None
            except (TypeError, ValueError):
                date = None

        return MessageModel(
            source=path,
            folder="",
            subject=subject,
            from_addr=_s(getattr(msg, "sender", "")),
            to_addrs=_s(getattr(msg, "to", "")),
            cc_addrs=_s(getattr(msg, "cc", "")),
            date=date,
            body_html=body_html,
            attachments=attachments,
            message_id=_s(getattr(msg, "messageId", None)).strip("<> ") or None,
            has_images=has_images or "<img" in body_html.lower(),
        )
    finally:
        close = getattr(msg, "close", None)
        if callable(close):
            close()


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        return _decode_text(value)
    return str(value)
